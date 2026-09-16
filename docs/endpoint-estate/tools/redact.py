#!/usr/bin/env python3
"""Redact every bearer credential before the workbook enters a repository.

A signed Power Automate URL is a credential: possession alone authorises invoking
the flow. So is an API key. The Master Reference redacts its own; the Harvest
Collector's raw error payloads redact nothing, and they carry both.

This first shipped matching `sig=` alone. That was too narrow: the same error
payloads also held four third-party API keys - OpenAI, OpenRouter, HuggingFace
and Google - which a signature-only scan walked straight past and which reached
a pushed commit before a wider sweep caught them. Match every credential SHAPE,
and keep each prefix so a reader can still tell what was there.

Usage:  python3 redact.py
"""
import os
import pathlib
import signal
signal.signal(signal.SIGPIPE, signal.SIG_DFL)  # writing into `| head` is not an error
import re
import openpyxl
from openpyxl.styles import Font

WB = os.environ.get('ESTATE_WORKBOOK',
    str(pathlib.Path(__file__).resolve().parent.parent / 'DGO_Unified_Endpoint_Estate_Merged.xlsx'))
# (pattern, replacement). The prefix survives; the secret does not.
RULES = [
    (re.compile(r'(sk-or-v1-)[A-Za-z0-9]{20,}'),                r'\1REDACTED'),
    (re.compile(r'(sk-ant-)[A-Za-z0-9_-]{20,}'),                r'\1REDACTED'),
    (re.compile(r'(sk-)(?!or-v1-|ant-|REDACTED)[A-Za-z0-9]{20,}'), r'\1REDACTED'),
    (re.compile(r'(hf_)(?!REDACTED)[A-Za-z0-9]{20,}'),          r'\1REDACTED'),
    (re.compile(r'(gh[pousr]_)(?!REDACTED)[A-Za-z0-9]{20,}'),   r'\1REDACTED'),
    (re.compile(r'(xox[baprs]-)(?!REDACTED)[A-Za-z0-9-]{10,}'), r'\1REDACTED'),
    (re.compile(r'(AKIA)(?!REDACTED)[0-9A-Z]{16}'),             r'\1REDACTED'),
    (re.compile(r'(AIza)(?!REDACTED)[0-9A-Za-z_-]{35}'),        r'\1REDACTED'),
    (re.compile(r'(eyJ[A-Za-z0-9_-]{10,}\.)eyJ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}'), r'\1REDACTED'),
    (re.compile(r'(sig=)(?!REDACTED)[A-Za-z0-9_\-%.]{8,}'),     r'\1REDACTED'),
    (re.compile(r'((?:AccountKey|SharedAccessKey|Password)=)(?!REDACTED)[^;"\s]{12,}'), r'\1REDACTED'),
]

wb = openpyxl.load_workbook(WB, data_only=False)
hits = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str):
                continue
            new = c.value
            for pat, repl in RULES:
                new = pat.sub(repl, new)
            if new != c.value:
                c.value = new
                hits.append(f'{ws.title}!{c.coordinate}')

idx = wb['00 Merge Index']
r = idx.max_row + 2
idx.cell(row=r, column=1, value='Redaction').font = Font(bold=True)
idx.cell(row=r + 1, column=1, value=(
    f'{len(hits)} cell(s) carried a bearer credential in raw harvest error payloads '
    f'({", ".join(hits)}). Each was replaced with its prefix plus "REDACTED", so a reader can '
    f'still tell what was there. Every credential found is live at its source until rotated - '
    f'removing it from a file revokes nothing. No other cell was changed.'))
idx.cell(row=r + 1, column=1).font = Font(italic=True)
wb.save(WB)
print(f'redacted {len(hits)} cell(s): {hits}')

check = openpyxl.load_workbook(WB, data_only=True)
live = [f'{ws.title}!{c.coordinate}'
        for ws in check.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str)
        for pat, _ in RULES if pat.search(c.value)]
print('live credentials remaining:', len(live), sorted(set(live)))
if live:
    raise SystemExit('redaction did not converge - a credential survived the pass')
