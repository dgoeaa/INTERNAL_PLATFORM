#!/usr/bin/env python3
"""Sheet-by-sheet and column-by-column audit: source vs merged."""
from pathlib import Path
import os
import pathlib
import signal
signal.signal(signal.SIGPIPE, signal.SIG_DFL)  # writing into `| head` is not an error
import openpyxl

SRC = Path(os.environ.get('ESTATE_SOURCE_DIR', '.'))
OUT = Path(os.environ.get('ESTATE_WORKBOOK',
    str(pathlib.Path(__file__).resolve().parent.parent / 'DGO_Unified_Endpoint_Estate_Merged.xlsx')))
SOURCES = [('1373e67d-DGO_Endpoint_Estate_Master_Reference.xlsx', ''),
           ('2a81b2c0-DGO_Flow_Harvest_Collector.xlsx', 'HC '),
           ('1d9b38f4-ECM_Simplified_Forensic_Audit.xlsx', 'FA '),
           ('70fd9e59-FlowIds_Endpoint_Database.xlsx', 'FI ')]

def header_row(ws):
    """The widest of the first six rows is the header row in these workbooks."""
    best, bi = [], 0
    for i, row in enumerate(ws.iter_rows(max_row=6, values_only=True), 1):
        vals = [v for v in row if v not in (None, '')]
        if len(vals) > len(best):
            best, bi = vals, i
    return bi, best

merged = openpyxl.load_workbook(OUT, data_only=False)
tot_sheets = tot_cols = missing = 0
for fname, prefix in SOURCES:
    wb = openpyxl.load_workbook(SRC / fname, data_only=False)
    print('=' * 92)
    print(fname[9:])
    for ws in wb.worksheets:
        tname = 'FI FlowIds Database' if prefix == 'FI ' else f'{prefix}{ws.title}'
        tws = merged[tname]
        hr, hdr = header_row(ws)
        _, thdr = header_row(tws)
        same = hdr == thdr
        # also compare the full column footprint, header or not
        def colset(w):
            return {c.column for r in w.iter_rows() for c in r if c.value is not None}
        cs, ts = colset(ws), colset(tws)
        ok = same and cs == ts
        missing += 0 if ok else 1
        tot_sheets += 1; tot_cols += len(hdr)
        print(f'  {ws.title:24s} -> {tname:24s} hdr_row={hr} cols={len(hdr):3d} '
              f'occupied_cols={len(cs):3d}  {"OK" if ok else "MISMATCH"}')
        if not ok:
            print('     source :', hdr)
            print('     merged :', thdr)
            print('     col diff:', cs ^ ts)
    wb.close()
print('=' * 92)
print(f'sheets audited: {tot_sheets}   header columns audited: {tot_cols}   mismatches: {missing}')
